"""TSRCT-PCB Analytical Tuning Logger V1.1.

Companion: Analytical-Tuning-uC-Firmware-V1.1.ino, Arduino Nano Every, 115200 baud.
Based on the supplied TSRCT_PCB_StepID_Autotune_v2_Logger (Hamish Finn / UTAS).
The requested filename Analyitical-Tuning-Logger.py is intentionally retained.

Install (Python 3.9+):
    pip install pyserial numpy pandas matplotlib openpyxl pillow
Optional nonlinear model cross-check:
    pip install scipy
Run after closing Arduino Serial Monitor, then press the board button:
    python Analyitical-Tuning-Logger.py --port COM5
Select iSIMC PID and its exported derivative filter:
    python Analyitical-Tuning-Logger.py --port COM5 --controller isimc-pid --filter-tau 3
Rebuild after interruption, or select the other tuning from the same capture:
    python Analyitical-Tuning-Logger.py --rebuild capture.csv --controller isimc-pid
    python Analyitical-Tuning-Logger.py --rebuild capture.csv --no-derivative-filter

The logger sends no commands. Closing it DOES NOT stop the heater experiment;
use the board button to abort. COMPLETE/ALARM/ABORTED telemetry ends acquisition.
CSV is the live accepted-data record; .session.jsonl preserves metadata, events
and rejected-frame diagnostics. Files flush every 5 s and fsync every 30 s.
Hard interruption can lose buffered data; rebuild skips incomplete records.

Only CRC-verified, structurally valid frames enter the live CSV. Genuine missing
sensor readings remain NaN in raw data, and are excluded from model fitting.
CRC-16 detects transport corruption; it is not sender authentication. CSV does
not retain frame CRC, so offline rebuild validates structure/ranges, not edits.

Identification: final pre-step baseline, 2% / 63.2% crossing solve, no delay floor.
Model times subtract the selected RTD's valid-readout age from the telemetry
clock; ADC conversion latency itself is not removed from the measured plant.
The selected heater's own RTD is always the fitted output; the other is optional
monitoring. The known rounded applied duty supplies the input scale. Completed
window identifiers distinguish the initial transition from steady scheduled duty.

Tuning convention follows main_final-2.pdf, section 5, equations (19)-(23):
  Kc = tau / (K * (tau_c + theta)); Ti = min(tau, 4*(tau_c + theta)).
  SIMC PI: tau_c=theta, Td=0. iSIMC PID: tau_c=theta/2, Td=theta/3.
  Series -> parallel: Kp=Kc*(1+Td/Ti), Ki=Kc/Ti, Kd=Kc*Td.
  Multiply all three by PWM_WINDOW_MS/100 for the nominal firmware's ms output.
  Units: Kp [ms/degC], Ki [ms/(degC*s)], Kd [ms*s/degC].
Primary rule reference: Grimholt and Skogestad, DYCOPS 2013:
  https://skoge.folk.ntnu.no/publications/2013/grimholt-dycops/0122.pdf
The finite derivative filter acts only on measurement derivative in nominal
V1.1. It changes the response; the manuscript comparison disabled that filter.
Default export is SIMC PI. One selected CH1_/CH2_ block replaces the matching
nominal settings. Export is blocked when required run/data/model gates fail;
candidate model diagnostics remain available. Export readiness is not a claim
of closed-loop or physical safety validation on a different thermal plant.
"""

from __future__ import annotations

import argparse
import binascii
import csv
import io
import json
import math
import os
import re
import signal
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

try:
    import serial
except ImportError:
    serial = None  # offline rebuild does not need a serial connection
try:
    from scipy.optimize import least_squares
except ImportError:
    least_squares = None

LOGGER_VERSION = "V1.1"
SCHEMA_VERSION = "analytical-1.1"
SERIAL_PORT = "COM5"
BAUD_RATE = 115200
SERIAL_TIMEOUT_S = 2.0
CSV_FLUSH_INTERVAL_S = 5.0
CSV_FSYNC_INTERVAL_S = 30.0
SILENCE_WARNING_S = 10.0
MAX_FRAME_BYTES = 4096
MAX_REJECTED_ROWS_IN_XLSX = 10_000

# Analysis policy, reported in the workbook. Missing rows are not interpolated.
MIN_COVERAGE = 0.95
REPORT_GAP_S = 1.5
MAX_BASELINE_GAP_S = 2.5
MAX_MODEL_GAP_S = 5.0
MAX_CROSSING_GAP_S = 1.5
MIN_RESOLVED_DELAY_S = 1.0
FINAL_AVERAGING_WINDOW_S = 600.0
FINAL_STEADY_MAX_ABS_SLOPE_C_PER_MIN = 0.02
MIN_STEP_SAMPLES = 60
MIN_DELTA_T_C = 0.5
MIN_MODEL_R2 = 0.98
ACTUATOR_TOLERANCE_PCT = 0.11
NLS_DELAY_ABS_WARNING_S = 2.0
NLS_DELAY_REL_WARNING = 0.5
DEFAULT_CONTROLLER = "simc-pi"
DEFAULT_DERIVATIVE_FILTER = True
DEFAULT_FILTER_TAU_S = 3.0

FIRMWARE_FIELDS = [
    "uptime_s", "experiment_s", "step_s", "ch1_C", "ch2_C", "test_channel",
    "monitor_other", "state", "alarm", "fault_flags", "ch1_valid", "ch2_valid",
    "ch1_rtd_age_ms", "ch2_rtd_age_ms", "ch1_rtd_sample_count", "ch2_rtd_sample_count",
    "ch1_rtd_failure_count", "ch2_rtd_failure_count", "ch1_rtd_fault_code", "ch2_rtd_fault_code",
    "commanded_duty_pct", "scheduled_gpio_duty_pct", "completed_command_duty_pct",
    "scheduled_window_us", "ssr_high_us", "window_sequence", "completed_window_sequence",
    "step_start_window_sequence", "step_start_uptime_ms", "run_id", "telemetry_seq",
    "rtd_state", "loop_max_us", "baseline_mean_C", "baseline_p2p_C", "baseline_count",
    "requested_duty_pct", "applied_duty_pct", "target_duration_s",
]
FLOAT_FIELDS = {
    "uptime_s", "experiment_s", "step_s", "ch1_C", "ch2_C", "commanded_duty_pct",
    "scheduled_gpio_duty_pct", "completed_command_duty_pct", "baseline_mean_C",
    "baseline_p2p_C", "requested_duty_pct", "applied_duty_pct", "target_duration_s",
}
INTEGER_FIELDS = set(FIRMWARE_FIELDS) - FLOAT_FIELDS
NAN_FIELDS = {"ch1_C", "ch2_C", "baseline_mean_C", "baseline_p2p_C"}
CSV_FIELDS = ["Timestamp", "Logger_elapsed_s"] + FIRMWARE_FIELDS
STATE_NAMES = {0: "IDLE", 1: "BASELINE", 2: "STEP_RUNNING", 3: "ALARM", 4: "COMPLETE", 5: "ABORTED"}
FAULT_BITS = {1: "CH1_OVERTEMP", 2: "CH2_OVERTEMP", 4: "CH1_RTD", 8: "CH2_RTD",
              16: "RTD_STALE", 32: "CONFIGURATION", 64: "COMMAND_TIMEOUT", 128: "WATCHDOG_RESET"}
EXPECTED_PARTS = len(FIRMWARE_FIELDS) + 1


def validate_row(row: dict[str, Any]) -> Optional[str]:
    for key in FIRMWARE_FIELDS:
        value = row[key]
        if key in INTEGER_FIELDS:
            if key == "test_channel":
                valid = value in (1, 2)
            elif key in ("alarm", "monitor_other", "ch1_valid", "ch2_valid"):
                valid = value in (0, 1)
            elif key == "state":
                valid = value in STATE_NAMES
            elif key == "rtd_state":
                valid = value in (0, 1, 2)
            elif key == "fault_flags" or key.endswith(("_failure_count", "_fault_code")):
                valid = 0 <= value <= 255
            elif key == "scheduled_window_us":
                valid = value == 1_000_000
            elif key == "ssr_high_us":
                valid = 0 <= value <= 1_000_000
            elif key == "baseline_count":
                valid = 0 <= value <= 120
            else:
                valid = 0 <= value <= 0xFFFFFFFF
            if not valid:
                return f"integer_range: {key}={value}"
        else:
            if key in NAN_FIELDS and math.isnan(value):
                continue
            if not math.isfinite(value):
                return f"nonfinite_value: {key}"
            if key in ("ch1_C", "ch2_C", "baseline_mean_C") and not -20 <= value <= 200:
                return f"temperature_range: {key}"
            if key == "baseline_p2p_C" and not 0 <= value <= 220:
                return "baseline_range"
            if key.endswith("duty_pct") and not 0 <= value <= 100:
                return f"duty_range: {key}"
            if key in ("experiment_s", "step_s") and value < 0 and value != -1:
                return f"time_range: {key}"
            if key in ("uptime_s", "target_duration_s") and value < 0:
                return f"time_range: {key}"
    if row["alarm"] != int(row["state"] == 3):
        return "inconsistent_alarm_flag"
    for ch in (1, 2):
        if row[f"ch{ch}_valid"] and (not math.isfinite(row[f"ch{ch}_C"]) or row[f"ch{ch}_rtd_age_ms"] > 1000):
            return f"inconsistent_sensor_validity: ch{ch}"
    if row["state"] in (0, 3, 4, 5) and any(row[k] != 0 for k in (
            "commanded_duty_pct", "scheduled_gpio_duty_pct", "completed_command_duty_pct", "ssr_high_us")):
        return "nonzero_output_in_off_state"
    if row["state"] == 2 and (row["step_s"] < 0 or row["step_start_window_sequence"] < 1):
        return "invalid_step_origin"
    if abs(row["scheduled_gpio_duty_pct"] - row["ssr_high_us"] / 10000) > 0.002:
        return "inconsistent_scheduled_duty"
    return None


def parse_data_line_diagnostic(line: str) -> tuple[Optional[dict[str, Any]], Optional[str]]:
    line = line.rstrip("\r\n")
    if len(line) > MAX_FRAME_BYTES:
        return None, "oversized_frame"
    if not line.startswith("D,"):
        return None, "not_D_prefix"
    try:
        raw = line.encode("ascii")
    except UnicodeEncodeError:
        return None, "non_ascii_frame"
    if any(c < 32 or c > 126 for c in raw):
        return None, "control_character_in_frame"
    payload, sep, checksum = line.rpartition("*")
    if not sep or not re.fullmatch(r"[0-9A-Fa-f]{4}", checksum):
        return None, "crc_framing_missing_or_invalid"
    if binascii.crc_hqx(payload.encode("ascii"), 0xFFFF) != int(checksum, 16):
        return None, "crc_mismatch"
    parts = payload.split(",")
    if len(parts) != EXPECTED_PARTS:
        return None, f"field_count: got {len(parts)}, expected {EXPECTED_PARTS}"
    row = {}
    for key, value in zip(FIRMWARE_FIELDS, parts[1:]):
        try:
            if key in INTEGER_FIELDS:
                if not re.fullmatch(r"[0-9]+", value):
                    raise ValueError
                row[key] = int(value)
            else:
                if value != "nan" and not re.fullmatch(r"-?[0-9]+(?:\.[0-9]+)?(?:[eE][+-]?[0-9]+)?", value):
                    raise ValueError
                row[key] = float(value)
        except (ValueError, OverflowError):
            return None, f"numeric_parse: {key}={value!r}"
    reason = validate_row(row)
    return (None if reason else row), reason


def parse_data_line(line: str) -> Optional[dict[str, Any]]:
    return parse_data_line_diagnostic(line)[0]


def journal_path_for(csv_path: Path) -> Path:
    return csv_path.with_suffix(".session.jsonl")


def read_capture(csv_path: Path) -> tuple[pd.DataFrame, int]:
    rows, skipped = [], 0
    with csv_path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames != CSV_FIELDS:
            raise ValueError("CSV header differs from the V1.1 analytical-tuning contract.")
        for record in reader:
            try:
                if None in record or any(record.get(k) is None for k in CSV_FIELDS):
                    raise ValueError("incomplete row")
                row = {key: int(record[key]) if key in INTEGER_FIELDS else float(record[key]) for key in FIRMWARE_FIELDS}
                if validate_row(row):
                    raise ValueError("invalid row")
                elapsed = float(record["Logger_elapsed_s"])
                if not math.isfinite(elapsed) or elapsed < 0:
                    raise ValueError("invalid receipt time")
                datetime.fromisoformat(record["Timestamp"])
                row.update(Timestamp=record["Timestamp"], Logger_elapsed_s=elapsed)
                rows.append(row)
            except (ValueError, TypeError, OverflowError):
                skipped += 1
    return pd.DataFrame(rows, columns=CSV_FIELDS), skipped


def finite_metadata(metadata: dict, key: str) -> float:
    try:
        value = float(metadata[key])
    except (ValueError, TypeError, KeyError):
        raise ValueError(f"Missing/invalid firmware metadata: {key}") from None
    if not math.isfinite(value):
        raise ValueError(f"Nonfinite firmware metadata: {key}")
    return value


def slope_per_min(t: np.ndarray, y: np.ndarray) -> float:
    if len(t) < 2 or np.ptp(t) <= 0:
        return float("nan")
    return float(np.polyfit(t - t[0], y, 1)[0] * 60)


def coverage(t: np.ndarray, expected_seconds: float) -> float:
    return min(1.0, len(t) / max(1.0, expected_seconds))


def max_gap(t: np.ndarray) -> float:
    return float(np.max(np.diff(t))) if len(t) > 1 else float("inf")


def crossing_time_linear(t: np.ndarray, y: np.ndarray, threshold: float) -> float:
    crossing = np.flatnonzero(y >= threshold)
    if len(crossing) == 0 or crossing[0] == 0:
        raise ValueError("Threshold crossing has no measured pre-crossing bracket.")
    i = int(crossing[0])
    dt = float(t[i] - t[i-1])
    if dt <= 0 or dt > MAX_CROSSING_GAP_S:
        raise ValueError(f"Threshold crossing spans a {dt:.3f} s gap; no interpolation permitted.")
    if not np.all(np.isfinite([t[i-1], t[i], y[i-1], y[i]])) or y[i] <= y[i-1]:
        raise ValueError("Invalid threshold crossing bracket.")
    return float(t[i-1] + (threshold-y[i-1])/(y[i]-y[i-1])*dt)


def fopdt_model(t: np.ndarray, T0: float, dT: float, theta: float, tau: float) -> np.ndarray:
    return T0 + dT * (1.0 - np.exp(-np.maximum(np.asarray(t, dtype=float)-theta, 0.0)/tau))


def fit_quality(measured: np.ndarray, fitted: np.ndarray) -> tuple[float, float]:
    ss_res = float(np.sum((measured-fitted)**2))
    ss_tot = float(np.sum((measured-np.mean(measured))**2))
    return (1-ss_res/ss_tot if ss_tot > 0 else float("nan")), float(np.sqrt(ss_res/len(measured)))


def compute_tunings(K: float, theta: float, tau: float, pwm_ms: float = 1000,
                    selected: str = DEFAULT_CONTROLLER, tau_c_factor: Optional[float] = None) -> list[dict]:
    """Complete unfiltered series-to-parallel conversion, then %-to-ms scaling."""
    if not all(math.isfinite(v) and v > 0 for v in (K, theta, tau, pwm_ms)):
        raise ValueError("Positive finite K, theta, tau and PWM window are required.")
    if selected not in ("simc-pi", "isimc-pid"):
        raise ValueError("Unknown controller selection")
    if tau_c_factor is not None and (not math.isfinite(tau_c_factor) or tau_c_factor <= 0):
        raise ValueError("tau-c factor must be positive and finite")
    rows = []
    for key, name, default_factor, td in (
            ("simc-pi", "SIMC PI", 1.0, 0.0),
            ("isimc-pid", "iSIMC PID", 0.5, theta/3.0)):
        factor = tau_c_factor if key == selected and tau_c_factor is not None else default_factor
        tc = factor*theta
        kc = tau/(K*(tc+theta))
        ti = min(tau, 4*(tc+theta))
        scale = pwm_ms/100.0
        rows.append({"key": key, "Controller": name, "tau_c_factor": factor, "tau_c_s": tc,
                     "Kc_pct_per_C": kc, "Ti_s": ti, "Td_s": td,
                     "Kp_ms_per_C": kc*(1+td/ti)*scale,
                     "Ki_ms_per_C_s": kc/ti*scale, "Kd_ms_s_per_C": kc*td*scale})
    if not all(math.isfinite(float(v)) for r in rows for k,v in r.items() if isinstance(v, (int,float))):
        raise ValueError("Nonfinite gain calculation")
    return rows


def nonlinear_crosscheck(plant: dict) -> dict:
    if least_squares is None:
        return {"success": False, "message": "SciPy unavailable; optional cross-check skipped"}
    t, y = plant["t"], plant["y"]
    duration = max(float(t[-1]), 1.0)
    lower = np.array([0.01, 0.0, 0.1])
    upper = np.array([max(plant["delta_T_C"]*3, 1), duration, max(duration*5, 10)])
    x0 = np.clip([plant["delta_T_C"], plant["theta_raw_s"], plant["tau_s"]], lower+1e-6, upper-1e-6)
    try:
        fitted = least_squares(lambda p: fopdt_model(t, plant["T0_C"], *p)-y,
                               x0=x0, bounds=(lower,upper), max_nfev=2000)
        dt, theta, tau = map(float, fitted.x)
        fit = fopdt_model(t, plant["T0_C"], dt, theta, tau)
        r2, rmse = fit_quality(y, fit)
        return {"success": bool(fitted.success), "message": str(fitted.message),
                "K_C_per_pct": dt/plant["duty_pct"], "theta_s": theta, "tau_s": tau,
                "R2": r2, "RMSE_C": rmse, "fit": fit}
    except Exception as exc:
        return {"success": False, "message": str(exc)}


def analyse(df: pd.DataFrame, metadata: dict, stats: dict, options: argparse.Namespace) -> dict:
    result = {"ready": False, "checks": [], "warnings": [], "baseline": None,
              "plant": None, "nls": None, "tunings": [], "channel": None}

    def check(name, passed, detail):
        result["checks"].append({"Check": name, "Pass": bool(passed), "Detail": str(detail)})

    check("Telemetry schema", metadata.get("schema") == SCHEMA_VERSION and stats.get("schema_verified", False),
          "Requires the matching firmware version/schema and field-order announcement")
    check("Board reset / ordering", not stats.get("board_resets", 0) and not stats.get("ordering_errors", 0),
          f"resets={stats.get('board_resets',0)}, ordering errors={stats.get('ordering_errors',0)}")
    if df.empty:
        check("Data available", False, "No accepted telemetry")
        return result
    run = df[df["run_id"] > 0].copy()
    one_run = run["run_id"].nunique() == 1
    check("Single experiment", one_run, f"{run['run_id'].nunique()} run identifiers")
    if not one_run:
        return result
    ch = int(run["test_channel"].iloc[0]); result["channel"] = ch
    other = 3-ch
    check("Channel identity", run["test_channel"].nunique() == 1 and run["monitor_other"].nunique() == 1,
          f"CH{ch} is fitted and receives exported gains")
    monitor = bool(run["monitor_other"].iloc[0])
    check("Time order", np.all(np.diff(run["uptime_s"].to_numpy(float)) > 0), "No mixed epochs or duplicate timestamps")
    terminal = run[run["state"] == 4]
    check("Successful terminal record", not terminal.empty and not run["state"].isin([3,5]).any(),
          "Requires CRC-valid COMPLETE telemetry and no alarm/abort")
    check("Fault-free run", bool((run["fault_flags"] == 0).all() and (run["alarm"] == 0).all()), "No latched firmware fault")
    try:
        duty = finite_metadata(metadata, "applied_duty_pct")
        requested = finite_metadata(metadata, "requested_duty_pct")
        pwm_ms = finite_metadata(metadata, "pwm_window_ms")
        duration = finite_metadata(metadata, "duration_s")
        baseline_seconds = finite_metadata(metadata, "baseline_window_samples")
        baseline_limit = finite_metadata(metadata, "baseline_max_p2p_C")
        baseline_min_s = finite_metadata(metadata, "baseline_min_s")
        configured_ch = finite_metadata(metadata, "test_channel")
        configured_monitor = finite_metadata(metadata, "monitor_other")
        if not (0 < duty <= 100 and 0 < requested <= 100 and pwm_ms == 1000 and duration > 0
                and 2 <= baseline_seconds <= 120 and baseline_limit > 0 and baseline_min_s >= baseline_seconds):
            raise ValueError("Unsupported/invalid firmware experiment configuration")
        rounded = math.floor(requested*pwm_ms/100+0.5)*100/pwm_ms
        check("Input/configuration agreement",
              abs(duty-rounded) <= 0.001 and configured_ch == ch and configured_monitor == int(monitor)
              and np.allclose(run["applied_duty_pct"], duty, atol=0.001, rtol=0)
              and np.allclose(run["requested_duty_pct"], requested, atol=0.0002, rtol=0)
              and np.allclose(run["target_duration_s"], duration, atol=0.001, rtol=0),
              f"Requested {requested:g}%, rounded/applied {duty:g}%, window {pwm_ms:g} ms")
    except ValueError as exc:
        check("Required experiment metadata", False, exc)
        return result
    step_all = run[run["state"] == 2].copy()
    if step_all.empty:
        check("Step data available", False, "No STEP_RUNNING records")
        return result
    origins = step_all["experiment_s"]-step_all["step_s"]
    onset = float(origins.median())
    check("Step boundary recorded", step_all["step_start_uptime_ms"].nunique() == 1
          and step_all["step_start_window_sequence"].nunique() == 1 and np.ptp(origins) <= 0.02
          and np.allclose(step_all["uptime_s"]-step_all["step_s"],
                          step_all["step_start_uptime_ms"]/1000, atol=0.02, rtol=0),
          f"Step begins at experiment t={onset:.3f} s; timestamp captured in ISR")
    check("Baseline minimum duration", onset >= baseline_min_s, f"{onset:.3f} / {baseline_min_s:g} s")
    completed = step_all[(step_all["completed_window_sequence"] >= step_all["step_start_window_sequence"])
                         & (step_all["step_start_window_sequence"] > 0)]
    duty_match = (not completed.empty and
                  np.all(np.abs(completed["completed_command_duty_pct"]-duty) <= ACTUATOR_TOLERANCE_PCT)
                  and np.all(np.abs(completed["scheduled_gpio_duty_pct"]-completed["completed_command_duty_pct"]) <= ACTUATOR_TOLERANCE_PCT)
                  and np.all(np.abs(step_all["commanded_duty_pct"]-duty) <= ACTUATOR_TOLERANCE_PCT))
    check("Completed-window duty agreement", duty_match,
          f"{len(completed)} corresponding complete windows checked; first transition excluded")

    healthy = (run[f"ch{ch}_valid"] == 1) & np.isfinite(run[f"ch{ch}_C"]) & (run["fault_flags"] == 0) & (run["alarm"] == 0)
    if monitor:
        healthy &= (run[f"ch{other}_valid"] == 1) & np.isfinite(run[f"ch{other}_C"])
    clean = run[healthy].copy()
    clean["measurement_experiment_s"] = clean["experiment_s"] - clean[f"ch{ch}_rtd_age_ms"]/1000.0
    clean["measurement_step_s"] = clean["step_s"] - clean[f"ch{ch}_rtd_age_ms"]/1000.0
    baseline = clean[(clean["state"] == 1) & (clean["measurement_experiment_s"] >= onset-baseline_seconds)
                     & (clean["measurement_experiment_s"] < onset)]
    if len(baseline) < 2:
        check("Baseline data", False, "Insufficient valid recent pre-step samples")
        return result
    bt = baseline["measurement_experiment_s"].to_numpy(float)
    by = baseline[f"ch{ch}_C"].to_numpy(float)
    bspan = float(bt[-1]-bt[0]); bp2p = float(np.ptp(by)); bcoverage = coverage(bt, baseline_seconds)
    baseline_good = (bspan >= baseline_seconds-2.5 and bcoverage >= MIN_COVERAGE
                     and max_gap(bt) <= MAX_BASELINE_GAP_S and onset-bt[-1] <= 1.5)
    check("Recent baseline coverage", baseline_good,
          f"{len(bt)} samples, span {bspan:.3f} s, coverage {bcoverage:.1%}, max gap {max_gap(bt):.3f} s")
    check("Baseline stability", bp2p <= baseline_limit, f"p-p {bp2p:.4f} <= {baseline_limit:g} degC")
    result["baseline"] = {"n": len(bt), "span_s": bspan, "coverage": bcoverage,
                          "mean_C": float(by.mean()), "std_C": float(by.std(ddof=1)),
                          "p2p_C": bp2p, "slope_C_per_min": slope_per_min(bt, by), "t": bt, "y": by}
    # The first step-state frame can still contain a pre-step sensor readout.
    step = clean[(clean["state"] == 2) & (clean["measurement_step_s"] >= 0)].copy()
    if len(step) < MIN_STEP_SAMPLES:
        check("Step data", False, f"{len(step)} valid samples; at least {MIN_STEP_SAMPLES} required")
        return result
    t, y = step["measurement_step_s"].to_numpy(float), step[f"ch{ch}_C"].to_numpy(float)
    if not np.all(np.diff(t) > 0):
        check("Step time monotonic", False, "No sorting or merging of conflicting timestamps")
        return result
    check("Full step coverage", t[0] <= 1.5 and t[-1] >= duration-1.5
          and coverage(t, duration) >= MIN_COVERAGE and max_gap(t) <= MAX_MODEL_GAP_S,
          f"{len(t)} samples; first/last {t[0]:.3f}/{t[-1]:.3f} s; target {duration:g} s; max gap {max_gap(t):.3f} s")
    check("Completed duration", not terminal.empty and bool((terminal["step_s"] >= duration-0.002).any()),
          "Terminal record must confirm the configured full duration")
    end_mask = t >= t[-1]-FINAL_AVERAGING_WINDOW_S
    ft, fy = t[end_mask], y[end_mask]
    fspan = float(ft[-1]-ft[0]); fslope = slope_per_min(ft,fy)
    check("Final averaging coverage", fspan >= FINAL_AVERAGING_WINDOW_S-2
          and coverage(ft, FINAL_AVERAGING_WINDOW_S) >= MIN_COVERAGE and max_gap(ft) <= MAX_MODEL_GAP_S,
          f"{len(ft)} samples spanning {fspan:.3f} / {FINAL_AVERAGING_WINDOW_S:g} s")
    check("Final steady-state slope", math.isfinite(fslope) and abs(fslope) <= FINAL_STEADY_MAX_ABS_SLOPE_C_PER_MIN,
          f"{fslope:+.5f} degC/min; limit +/-{FINAL_STEADY_MAX_ABS_SLOPE_C_PER_MIN:g}")
    t0, tf = float(by.mean()), float(fy.mean())
    delta = tf-t0
    check("Positive thermal step", delta >= MIN_DELTA_T_C, f"Delta T={delta:.4f} degC")
    if delta < MIN_DELTA_T_C:
        return result
    try:
        t2 = crossing_time_linear(t,y,t0+0.02*delta)
        t63 = crossing_time_linear(t,y,t0+0.632*delta)
        a2, a63 = -math.log(0.98), -math.log(1-0.632)
        tau = (t63-t2)/(a63-a2)
        theta = t2-a2*tau  # retained exactly; no hidden 1 s clamp
        if not all(math.isfinite(v) for v in (tau,theta)) or tau <= 0:
            raise ValueError("Nonphysical two-crossing solution")
        check("Crossing brackets", True, f"t2={t2:.4f} s, t63.2={t63:.4f} s")
    except ValueError as exc:
        check("Crossing brackets", False, exc)
        return result
    fit = fopdt_model(t,t0,delta,theta,tau)
    r2, rmse = fit_quality(y,fit)
    K = delta/duty
    check("Resolved positive delay", theta >= MIN_RESOLVED_DELAY_S,
          f"theta_raw={theta:.6f} s; minimum resolved delay {MIN_RESOLVED_DELAY_S:g} s; no floor applied")
    check("Model fit", math.isfinite(r2) and r2 >= MIN_MODEL_R2, f"R2={r2:.6f}, RMSE={rmse:.5f} degC")
    plant = {"K_C_per_pct": K, "theta_raw_s": theta, "tau_s": tau, "T0_C": t0, "Tf_C": tf,
             "delta_T_C": delta, "duty_pct": duty, "t2_s": t2, "t63_s": t63, "R2": r2, "RMSE_C": rmse,
             "final_slope_C_per_min": fslope, "final_span_s": fspan, "final_samples": len(ft),
             "t": t, "y": y, "fit": fit}
    result["plant"] = plant
    if not options.no_nls:
        result["nls"] = nonlinear_crosscheck(plant)
        nls = result["nls"]
        if nls.get("success"):
            diff = abs(nls["theta_s"]-theta)
            if diff > max(NLS_DELAY_ABS_WARNING_S,NLS_DELAY_REL_WARNING*max(theta,0)):
                result["warnings"].append(f"Delay estimates disagree: two-crossing {theta:.3f} s, nonlinear {nls['theta_s']:.3f} s. Review the early response.")
        else:
            result["warnings"].append(nls["message"])
    if theta > 0:
        try:
            result["tunings"] = compute_tunings(K,theta,tau,pwm_ms,options.controller,options.tau_c_factor)
        except ValueError as exc:
            check("Finite controller gains", False, exc)
    else:
        check("Finite controller gains", False, "Positive delay required; gain calculation suppressed")
    result["ready"] = all(c["Pass"] for c in result["checks"]) and bool(result["tunings"])
    return result


def arduino_lines(analysis: dict, options: argparse.Namespace) -> list[str]:
    if not analysis["ready"]:
        return ["// EXPORT BLOCKED: identification is not qualified."] + [
            f"// {c['Check']}: {c['Detail']}" for c in analysis["checks"] if not c["Pass"]]
    chosen = next(r for r in analysis["tunings"] if r["key"] == options.controller)
    ch, plant = analysis["channel"], analysis["plant"]
    lines = [
        f"// CH{ch} {chosen['Controller']} from qualified open-loop identification.",
        f"// K={plant['K_C_per_pct']:.9g} C/%; theta={plant['theta_raw_s']:.9g} s; tau={plant['tau_s']:.9g} s.",
        f"// tau_c={chosen['tau_c_s']:.9g} s; Ti={chosen['Ti_s']:.9g} s; Td={chosen['Td_s']:.9g} s.",
        "// Replace these matching lines in the V1.1 nominal firmware USER SETTINGS.",
        "// Kp [ms/C], Ki [ms/(C*s)], Kd [ms*s/C]; 1000 ms SSR window, 1 s control sample.",
        f"const float CH{ch}_KP = {chosen['Kp_ms_per_C']:.9f}f;",
        f"const float CH{ch}_KI = {chosen['Ki_ms_per_C_s']:.9f}f;",
        f"const float CH{ch}_KD = {chosen['Kd_ms_s_per_C']:.9f}f;",
        f"const bool CH{ch}_DERIVATIVE_FILTER_ENABLED = {'true' if options.derivative_filter else 'false'};",
        f"const float CH{ch}_DERIVATIVE_TAU_S = {options.filter_tau:.6f}f;",
        f"// Filter alpha=Ts/(tau_f+Ts)={1/(options.filter_tau+1):.9f}; Ts=1 s.",
        "// Derivative is on measurement. Kd=0 disables derivative action for PI.",
        "// The manuscript PID comparison used filtering disabled; enabled filtering changes response.",
        "// Keep the nominal target, channel mode and protection settings appropriate to this plant.",
    ]
    lines += [f"// Review note: {warning}" for warning in analysis["warnings"]]
    return lines


def gap_trace(t: np.ndarray, y: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    """Insert blank plotting points across gaps without adding measurements."""
    t, y = np.asarray(t,dtype=float), np.asarray(y,dtype=float)
    indices = np.flatnonzero((np.diff(t) > REPORT_GAP_S) | (np.diff(t) <= 0))+1
    return np.insert(t,indices,np.nan), np.insert(y,indices,np.nan)


def figure_buffer(fig) -> io.BytesIO:
    output = io.BytesIO()
    fig.tight_layout()
    fig.savefig(output,format="png",dpi=150,bbox_inches="tight")
    plt.close(fig); output.seek(0)
    return output


def make_plots(df: pd.DataFrame, analysis: dict) -> list[tuple[str, io.BytesIO]]:
    plots = []
    if not df.empty:
        fig, ax = plt.subplots(figsize=(10.2,4.8))
        t = df["Logger_elapsed_s"].to_numpy(float)
        for ch,color in ((1,"#1764a2"),(2,"#d57820")):
            y = df[f"ch{ch}_C"].where(df[f"ch{ch}_valid"] == 1).to_numpy(float)
            xg,yg = gap_trace(t,y)
            ax.plot(xg/60,yg,lw=0.9,color=color,label=f"CH{ch} temperature")
        duty_ax = ax.twinx()
        for key,label,style in (("commanded_duty_pct","Current command",":"),
                                ("scheduled_gpio_duty_pct","Completed scheduled duty","--")):
            xg,yg = gap_trace(t,df[key].to_numpy(float))
            duty_ax.plot(xg/60,yg,style,lw=0.8,label=label,alpha=0.7)
        duty_ax.set_ylabel("Duty (%)"); duty_ax.set_ylim(-1,101)
        ax.set(xlabel="Logger elapsed time (min)",ylabel="Temperature (degC)",title="Analytical tuning: complete acquisition")
        handles,labels=ax.get_legend_handles_labels(); h2,l2=duty_ax.get_legend_handles_labels()
        ax.legend(handles+h2,labels+l2,fontsize=8,loc="best"); ax.grid(alpha=0.25)
        plots.append(("Complete capture; gaps remain blank",figure_buffer(fig)))
    baseline = analysis["baseline"]
    if baseline:
        fig,ax=plt.subplots(figsize=(10.2,4.4))
        xg,yg=gap_trace(baseline["t"],baseline["y"])
        ax.plot(xg-baseline["t"][0],yg,lw=1,label=f"CH{analysis['channel']}")
        ax.axhline(baseline["mean_C"],ls="--",lw=1,label="Baseline mean")
        ax.set(xlabel="Time within final baseline window (s)",ylabel="Temperature (degC)",
               title=f"Pre-step baseline: p-p={baseline['p2p_C']:.4f} degC; coverage={baseline['coverage']:.1%}")
        ax.grid(alpha=0.25); ax.legend(fontsize=8)
        plots.append(("Recent baseline qualification",figure_buffer(fig)))
    plant = analysis["plant"]
    if plant:
        fig,(ax,res)=plt.subplots(2,1,figsize=(10.2,6.3),sharex=True,gridspec_kw={"height_ratios":[2,1]})
        t,y=plant["t"],plant["y"]
        xg,yg=gap_trace(t,y)
        ax.plot(xg/60,yg,lw=0.9,label="Measured selected RTD")
        ax.plot(t/60,plant["fit"],"--",lw=1.2,label="Two-crossing FOPDT")
        nls=analysis["nls"]
        if nls and nls.get("success"):
            ax.plot(t/60,nls["fit"],":",lw=1,label="Optional nonlinear cross-check")
        ax.set(ylabel="Temperature (degC)",title=f"FOPDT: K={plant['K_C_per_pct']:.5g} degC/%; theta={plant['theta_raw_s']:.3f} s; tau={plant['tau_s']:.3f} s")
        ax.grid(alpha=0.25); ax.legend(fontsize=8)
        xg,yg=gap_trace(t,y-plant["fit"])
        res.plot(xg/60,yg,lw=0.7); res.axhline(0,ls="--",lw=0.8)
        res.set(xlabel="Time after applied step (min)",ylabel="Residual (degC)",
                title=f"R2={plant['R2']:.6f}; RMSE={plant['RMSE_C']:.5f} degC")
        res.grid(alpha=0.25)
        plots.append(("FOPDT fit and residuals",figure_buffer(fig)))
        fig,ax=plt.subplots(figsize=(10.2,4.4))
        early=t <= max(plant["t2_s"]*2, plant["theta_raw_s"]+0.1*plant["tau_s"], 20)
        xg,yg=gap_trace(t[early],y[early])
        ax.plot(xg,yg,".-",ms=2,lw=0.8,label="Measured early response")
        ax.plot(t[early],plant["fit"][early],"--",lw=1,label="FOPDT")
        ax.axvline(plant["theta_raw_s"],ls=":",label="Raw identified delay")
        ax.axvline(plant["t2_s"],ls="--",label="2% crossing")
        ax.axhline(plant["T0_C"]+0.02*plant["delta_T_C"],ls=":",lw=0.8)
        ax.set(xlabel="Time after applied step (s)",ylabel="Temperature (degC)",title="Early response and delay estimate; no delay floor applied")
        ax.grid(alpha=0.25); ax.legend(fontsize=8)
        plots.append(("Early response / delay resolution",figure_buffer(fig)))
    return plots


def excel_value(value):
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value,float) and not math.isfinite(value):
        return None
    if isinstance(value,str):
        return ILLEGAL_CHARACTERS_RE.sub("",value)[:32767]
    return value


def add_sheet(wb, name: str, headers: list[str], rows) -> Any:
    ws = wb.create_sheet(name); ws.append(headers)
    for record in rows:
        ws.append([excel_value(value) for value in record])
    for cell in ws[1]:
        cell.font=Font(bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="234D6A")
        cell.alignment=Alignment(vertical="center",wrap_text=True)
    ws.row_dimensions[1].height=30
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    for i,key in enumerate(headers,1):
        ws.column_dimensions[get_column_letter(i)].width=min(36,max(14,len(key)+2))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value,str):
                cell.data_type="s"  # startup separators never become Excel formulas
                cell.alignment=Alignment(vertical="top",wrap_text=name != "Timeseries")
            elif isinstance(cell.value,float):
                cell.number_format="0.000000"
    return ws


def write_workbook(path: Path, csv_path: Path, df: pd.DataFrame, metadata: dict, stats: dict,
                   events: list[dict], rejected: list[dict], analysis: dict,
                   options: argparse.Namespace, termination: str) -> None:
    wb=openpyxl.Workbook(); wb.remove(wb.active)
    ready="READY FOR GAIN EXPORT" if analysis["ready"] else "EXPORT BLOCKED — see Readiness"
    summary=[("Status",ready),("Logger version",LOGGER_VERSION),("Firmware",metadata.get("firmware_version","unknown")),
             ("Termination",termination),("Generated",datetime.now().isoformat(timespec="seconds")),
             ("Selected controller",options.controller),("Identified channel",analysis["channel"]),
             ("Derivative filter enabled",options.derivative_filter),("Derivative filter tau (s)",options.filter_tau),
             ("Derivative filter alpha at Ts=1 s",1/(1+options.filter_tau)),
             ("Accepted telemetry rows",len(df)),("Rejected frames",stats.get("rejected",0)),
             ("CSV",str(csv_path.resolve())),("Journal",str(journal_path_for(csv_path).resolve())),
             ("GPIO interpretation","Scheduled ISR duty; no independent GPIO, SSR or heater-power measurement"),
             ("Estimator","2% / 63.2% baseline-referenced crossing solve; raw delay retained"),
             ("Model time basis","Selected RTD valid-readout timestamp relative to the recorded actuation boundary; ADC conversion latency retained"),
             ("PI convention","tau_c=theta; Td=0; Ti=min(tau,4*(tau_c+theta))"),
             ("PID convention","tau_c=theta/2; Td=theta/3; complete series-to-parallel conversion"),
             ("Custom tau_c factor",options.tau_c_factor if options.tau_c_factor is not None else "Defaults above"),
             ("Filter convention","Filter is applied to measurement derivative in nominal V1.1; manuscript comparison was unfiltered"),
             ("Scope","Qualified identification provides candidate gains for this plant; verify closed-loop behaviour on the bench")]
    summary += [("Review note", warning) for warning in analysis["warnings"]]
    ws=add_sheet(wb,"Run Summary",["Item","Value"],summary)
    ws.column_dimensions["A"].width=36; ws.column_dimensions["B"].width=105
    for row in range(2,ws.max_row+1): ws.row_dimensions[row].height=30
    ws=add_sheet(wb,"Readiness",["Check","Pass","Detail"],
                 [(c["Check"],"PASS" if c["Pass"] else "FAIL",c["Detail"]) for c in analysis["checks"]])
    ws.column_dimensions["A"].width=34; ws.column_dimensions["C"].width=100
    for row in range(2,ws.max_row+1): ws.row_dimensions[row].height=30
    tuning_headers=["Controller","tau_c_factor","tau_c_s","Kc_pct_per_C","Ti_s","Td_s",
                    "Kp_ms_per_C","Ki_ms_per_C_s","Kd_ms_s_per_C"]
    qualification="Qualified candidate" if analysis["ready"] else "DIAGNOSTIC ONLY — EXPORT BLOCKED"
    add_sheet(wb,"Controller Settings",["Status"]+tuning_headers,
              [[qualification]+[r[k] for k in tuning_headers] for r in analysis["tunings"]])
    ws=add_sheet(wb,"Arduino Ready",["Selected nominal V1.1 settings"],[(line,) for line in arduino_lines(analysis,options)])
    ws.column_dimensions["A"].width=125
    for row in ws.iter_rows(min_row=2):
        row[0].font=Font(name="Consolas",size=10); ws.row_dimensions[row[0].row].height=30
    parameters=[]
    for estimator, data in (("Baseline",analysis["baseline"]),("Two-crossing",analysis["plant"]),("Nonlinear cross-check",analysis["nls"])):
        if data:
            parameters.extend((estimator,k,v) for k,v in data.items() if not isinstance(v,np.ndarray))
    ws=add_sheet(wb,"Model Diagnostics",["Estimator","Parameter (units in name)","Value"],parameters)
    ws.column_dimensions["C"].width=65
    ws=add_sheet(wb,"Configuration",["Key","Value"],list(metadata.items()))
    ws.column_dimensions["A"].width=38; ws.column_dimensions["B"].width=70
    policy={"Minimum coverage":MIN_COVERAGE,"Reported gap threshold (s)":REPORT_GAP_S,
            "Maximum baseline gap (s)":MAX_BASELINE_GAP_S,"Maximum model gap (s)":MAX_MODEL_GAP_S,
            "Maximum crossing bracket gap (s)":MAX_CROSSING_GAP_S,"Minimum resolved delay (s)":MIN_RESOLVED_DELAY_S,
            "Final averaging window (s)":FINAL_AVERAGING_WINDOW_S,"Final slope limit (C/min)":FINAL_STEADY_MAX_ABS_SLOPE_C_PER_MIN,
            "Minimum R2":MIN_MODEL_R2,"Duty tolerance (% points)":ACTUATOR_TOLERANCE_PCT}
    integrity=list(stats.items())+list(policy.items())
    if not df.empty:
        for ch in (1,2): integrity.append((f"CH{ch} sensor-invalid rows",int((df[f"ch{ch}_valid"] == 0).sum())))
        times=df["uptime_s"].to_numpy(float)
        integrity += [("Intervals above 1.5 s",int(np.sum(np.diff(times)>REPORT_GAP_S))),
                      ("Maximum telemetry interval (s)",max_gap(times))]
    add_sheet(wb,"Data Integrity",["Metric","Value"],integrity)
    ws=add_sheet(wb,"Events",["Timestamp","Event"],[(e.get("Timestamp",""),e.get("Event","")) for e in events])
    ws.column_dimensions["B"].width=110
    ws=add_sheet(wb,"Rejected Rows",["Timestamp","Reason","Raw Line (escaped)"],
                 [(r.get("Timestamp",""),r.get("Reason",""),r.get("Raw Line","")) for r in rejected[:MAX_REJECTED_ROWS_IN_XLSX]])
    ws.column_dimensions["B"].width=45; ws.column_dimensions["C"].width=115
    ws=wb.create_sheet("Plots"); ws.sheet_view.showGridLines=False
    row_anchor=1
    for title, buffer in make_plots(df,analysis):
        ws.cell(row=row_anchor,column=1,value=title).font=Font(bold=True,size=13)
        img=XLImage(buffer)
        factor=min(1.0,950.0/img.width)
        img.width=round(img.width*factor); img.height=round(img.height*factor)
        ws.add_image(img,f"A{row_anchor+1}")
        next_anchor=row_anchor+math.ceil(img.height/20)+4
        for row in range(row_anchor,next_anchor+1): ws.row_dimensions[row].height=15
        ws.row_dimensions[row_anchor].height=24
        row_anchor=next_anchor
    ws=add_sheet(wb,"Timeseries",CSV_FIELDS,df[CSV_FIELDS].itertuples(index=False,name=None))
    for row in ws.iter_rows(min_row=2):
        for key,cell in zip(CSV_FIELDS,row):
            if key in INTEGER_FIELDS: cell.number_format="0"
            elif key.endswith("_C"): cell.number_format="0.0000"
            elif key != "Timestamp": cell.number_format="0.000"
    order=["Run Summary","Readiness","Controller Settings","Arduino Ready","Plots","Model Diagnostics",
           "Configuration","Data Integrity","Events","Rejected Rows","Timeseries"]
    wb._sheets=[wb[name] for name in order]
    temporary=Path(str(path)+".tmp.xlsx")
    try:
        wb.save(temporary); os.replace(temporary,path)
    finally:
        if temporary.exists(): temporary.unlink()


def unique_paths(output: Optional[Path] = None) -> tuple[Path,Path]:
    if output is None:
        output=Path(f"Analytical-Tuning-V1.1_{datetime.now():%Y-%m-%d_%H-%M-%S}.xlsx")
    if output.suffix.lower() != ".xlsx": output=Path(str(output)+".xlsx")
    output.parent.mkdir(parents=True,exist_ok=True)
    original=output; number=1
    while output.exists() or output.with_suffix(".csv").exists() or journal_path_for(output.with_suffix(".csv")).exists():
        output=original.with_name(f"{original.stem}_{number:03d}.xlsx"); number+=1
    return output.with_suffix(".csv"),output


def build_report(csv_path: Path, xlsx_path: Path, metadata: dict, stats: dict,
                 events: list[dict], rejected: list[dict], options: argparse.Namespace,
                 termination: str) -> dict:
    df, skipped=read_capture(csv_path)
    stats=dict(stats); stats["accepted_rows"]=len(df); stats["invalid_CSV_rows_skipped"]=skipped
    if skipped:
        events=list(events)+[{"Timestamp":datetime.now().isoformat(timespec="seconds"),
                              "Event":f"REBUILD: skipped {skipped} invalid/incomplete CSV rows"}]
    analysis=analyse(df,metadata,stats,options)
    print("[ANALYSIS] "+("READY FOR GAIN EXPORT" if analysis["ready"] else "EXPORT BLOCKED"))
    for check in analysis["checks"]:
        if not check["Pass"]: print(f"  {check['Check']}: {check['Detail']}")
    for warning in analysis["warnings"]: print(f"[REVIEW] {warning}")
    for tuning in analysis["tunings"]:
        print(f"  {tuning['Controller']}: Kp={tuning['Kp_ms_per_C']:.6f}, Ki={tuning['Ki_ms_per_C_s']:.9f}, Kd={tuning['Kd_ms_s_per_C']:.6f}")
    write_workbook(xlsx_path,csv_path,df,metadata,stats,events,rejected,analysis,options,termination)
    print(f"[SAVED] {xlsx_path.resolve()}")
    return analysis


def run_logger(options: argparse.Namespace) -> int:
    if serial is None:
        raise SystemExit("pyserial is required for acquisition: pip install pyserial")
    csv_path,xlsx_path=unique_paths(options.output)
    metadata={}; events=[]; rejected=[]
    stats={"accepted_rows":0,"rejected":0,"crc_rejected":0,"missing_sequences":0,
           "ordering_errors":0,"board_resets":0,"schema_verified":False}
    termination="serial_interruption"
    stop=False; exit_code=0; previous_row=None
    csv_handle=journal_handle=ser=None
    pending=bytearray(); oversized=False
    start=time.monotonic(); last_valid=last_flush=last_sync=last_print=start
    last_silence_warning=start-SILENCE_WARNING_S

    def record(kind: str, value: dict):
        if journal_handle is not None:
            journal_handle.write(json.dumps({"kind":kind,"record":value},ensure_ascii=True)+"\n")

    def event(message: str):
        entry={"Timestamp":datetime.now().isoformat(timespec="milliseconds"),"Event":message}
        events.append(entry); record("event",entry)

    def checkpoint():
        record("session",{"metadata":metadata,"stats":stats,"termination":termination})

    def reject(line: str, reason: str):
        stats["rejected"]+=1
        if reason.startswith("crc_"): stats["crc_rejected"]+=1
        entry={"Timestamp":datetime.now().isoformat(timespec="milliseconds"),"Reason":reason,"Raw Line":ascii(line)}
        if len(rejected)<MAX_REJECTED_ROWS_IN_XLSX: rejected.append(entry)
        record("rejected",entry)
        if stats["rejected"]<=10 or stats["rejected"]%100==0:
            print(f"[REJECT] #{stats['rejected']} {reason}: {ascii(line)}")

    def request_stop(_sig=None,_frame=None):
        nonlocal stop,termination
        stop=True; termination="user_interrupt"

    print(f"Analytical Tuning Logger {LOGGER_VERSION}; {options.port} @ {BAUD_RATE}")
    print(f"[CSV] {csv_path.resolve()}")
    print("[LOG] Start/abort the experiment using the board button. Closing Python does not stop heating.")
    previous_handler=signal.signal(signal.SIGINT,request_stop)
    try:
        csv_handle=csv_path.open("x",newline="",encoding="utf-8")
        journal_handle=journal_path_for(csv_path).open("x",encoding="utf-8")
        writer=csv.DictWriter(csv_handle,fieldnames=CSV_FIELDS); writer.writeheader()
        checkpoint()
        for handle in (csv_handle,journal_handle): handle.flush(); os.fsync(handle.fileno())
        ser=serial.Serial(options.port,BAUD_RATE,timeout=SERIAL_TIMEOUT_S)
        while not stop:
            raw=ser.readline(MAX_FRAME_BYTES+1)
            now=time.monotonic()
            if now-last_flush>=CSV_FLUSH_INTERVAL_S:
                checkpoint()
                for handle in (csv_handle,journal_handle): handle.flush()
                last_flush=now
            if now-last_sync>=CSV_FSYNC_INTERVAL_S:
                for handle in (csv_handle,journal_handle): handle.flush(); os.fsync(handle.fileno())
                last_sync=now
            if now-last_valid>=SILENCE_WARNING_S and now-last_silence_warning>=30:
                message=f"NO_VALID_TELEMETRY for {now-last_valid:.0f} s; board operation is independent. Check connection; local button aborts."
                print(f"[WARN] {message}"); event(message); last_silence_warning=now
            if not raw: continue
            if len(pending)+len(raw)>MAX_FRAME_BYTES: oversized=True
            pending.extend(raw[:max(0,MAX_FRAME_BYTES-len(pending))])
            if not raw.endswith(b"\n"): continue
            line=bytes(pending).decode("ascii",errors="replace").rstrip("\r\n"); pending.clear()
            if oversized:
                reject(line,"oversized_frame"); oversized=False; continue
            if not line: continue
            if any(ord(c)<32 or ord(c)>126 for c in line):
                reject(line,"non_ascii_or_control_character"); continue
            if line.startswith("# FIELDS:"):
                fields=line.split(":",1)[1].strip().split(",")
                if fields != FIRMWARE_FIELDS:
                    termination="firmware_field_contract_mismatch"; event(termination); exit_code=1; break
                stats["schema_verified"]=True; checkpoint(); continue
            if line.startswith("# CONFIG:"):
                for pair in line.split(":",1)[1].strip().split(","):
                    if "=" not in pair: continue
                    key,value=pair.split("=",1); key=key.strip(); value=value.strip()
                    if key in metadata and metadata[key] != value:
                        event(f"CONFIGURATION_CHANGE: {key}: {metadata[key]} -> {value}")
                        # A changed experiment configuration cannot be merged safely.
                        if stats["accepted_rows"] and key != "reset_cause":
                            stats["ordering_errors"]+=1
                    metadata[key]=value
                checkpoint()
                if metadata.get("schema",SCHEMA_VERSION)!=SCHEMA_VERSION:
                    termination="firmware_schema_mismatch"; event(termination); exit_code=1; break
                continue
            if line.startswith("#"):
                event(line.removeprefix("# EVENT:").strip() if line.startswith("# EVENT:") else line.lstrip("# "))
                # Events inform the record. Only CRC-valid terminal D rows end a run.
                continue
            row,reason=parse_data_line_diagnostic(line)
            if row is None:
                reject(line,reason or "unknown_parser_failure"); continue
            if previous_row is not None:
                delta=(row["telemetry_seq"]-previous_row["telemetry_seq"]) & 0xFFFFFFFF
                uptime_wrap=previous_row["uptime_s"]>4_294_000 and row["uptime_s"]<1000
                if row["uptime_s"]<previous_row["uptime_s"] and not uptime_wrap:
                    stats["board_resets"]+=1; event("BOARD_RESET: capture stopped to keep epochs separate")
                    termination="board_reset"; exit_code=1; stop=True
                elif delta==0 or delta>=0x80000000:
                    stats["ordering_errors"]+=1; reject(line,"duplicate_or_out_of_order_sequence"); continue
                elif delta>1:
                    stats["missing_sequences"]+=delta-1; event(f"TELEMETRY_GAP: {delta-1} missing emitted frame(s)")
            previous_row=row
            stamped={"Timestamp":datetime.now().isoformat(timespec="milliseconds"),
                     "Logger_elapsed_s":round(now-start,6),**row}
            writer.writerow(stamped); stats["accepted_rows"]+=1; last_valid=now
            if now-last_print>=5:
                ch=row["test_channel"]
                print(f"[DATA] {STATE_NAMES[row['state']]} CH{ch}: {row[f'ch{ch}_C']:.4f} C, step={row['step_s']:.1f} s, duty={row['commanded_duty_pct']:.2f}%")
                last_print=now
            if row["state"] in (3,4,5):
                termination=STATE_NAMES[row["state"]]; stop=True
    except (serial.SerialException,OSError) as exc:
        termination=f"acquisition_error: {exc}"; exit_code=1; print(f"[ERROR] {termination}")
    except KeyboardInterrupt:
        termination="user_interrupt"
    finally:
        signal.signal(signal.SIGINT,previous_handler)
        if ser is not None:
            try: ser.close()
            except (OSError,serial.SerialException): pass
        try:
            if pending: reject(bytes(pending).decode("ascii",errors="replace"),"incomplete_frame_at_shutdown")
            checkpoint()
        except OSError as exc:
            print(f"[ERROR] Journal finalisation: {exc}"); exit_code=1
        for handle in (csv_handle,journal_handle):
            if handle is not None:
                try: handle.flush(); os.fsync(handle.fileno())
                except OSError as exc:
                    print(f"[ERROR] Saving buffered data: {exc}"); exit_code=1
                finally:
                    try: handle.close()
                    except OSError: exit_code=1
    if csv_path.exists():
        try: build_report(csv_path,xlsx_path,metadata,stats,events,rejected,options,termination)
        except Exception as exc:
            print(f"[ERROR] Report generation failed: {exc}")
            print(f"[RECOVERY] Existing CSV/journal: {csv_path.resolve()}; use --rebuild")
            exit_code=1
    return exit_code


def rebuild(options: argparse.Namespace) -> int:
    csv_path=options.rebuild
    metadata={}; stats={}; events=[]; rejected=[]; termination="offline_rebuild"
    journal_path=journal_path_for(csv_path)
    if journal_path.exists():
        with journal_path.open(encoding="utf-8") as handle:
            for line in handle:
                try:
                    entry=json.loads(line); kind=entry["kind"]; record=entry["record"]
                    if kind=="session":
                        metadata=record.get("metadata",{}); stats=record.get("stats",{})
                        termination="offline_rebuild: "+record.get("termination","unknown")
                    elif kind=="event": events.append(record)
                    elif kind=="rejected" and len(rejected)<MAX_REJECTED_ROWS_IN_XLSX: rejected.append(record)
                except (ValueError,TypeError,KeyError): continue
    else:
        print("[WARN] Journal missing. Raw-data plots remain available; required metadata cannot be assumed.")
    output=options.output or csv_path.with_suffix(".rebuilt.xlsx")
    _unused,xlsx_path=unique_paths(output)
    build_report(csv_path,xlsx_path,metadata,stats,events,rejected,options,termination)
    return 0


def parse_args(argv=None) -> argparse.Namespace:
    parser=argparse.ArgumentParser(description=f"TSRCT-PCB Analytical Tuning Logger {LOGGER_VERSION}")
    parser.add_argument("--port",default=SERIAL_PORT)
    parser.add_argument("--output",type=Path,help="XLSX path; existing captures are never overwritten")
    parser.add_argument("--rebuild",type=Path,metavar="CSV",help="offline workbook recovery/retuning from a V1.1 CSV and adjacent journal")
    parser.add_argument("--controller",choices=("simc-pi","isimc-pid"),default=DEFAULT_CONTROLLER)
    parser.add_argument("--tau-c-factor",type=float,help="override tau_c/theta for the selected controller only")
    parser.add_argument("--filter-tau",type=float,default=DEFAULT_FILTER_TAU_S,help="exported derivative filter time constant in seconds")
    parser.add_argument("--no-derivative-filter",dest="derivative_filter",action="store_false",default=DEFAULT_DERIVATIVE_FILTER)
    parser.add_argument("--no-nls",action="store_true",help="skip optional nonlinear model cross-check")
    options=parser.parse_args(argv)
    if not math.isfinite(options.filter_tau) or options.filter_tau<0:
        parser.error("--filter-tau must be finite and non-negative")
    if options.tau_c_factor is not None and (not math.isfinite(options.tau_c_factor) or options.tau_c_factor<=0):
        parser.error("--tau-c-factor must be finite and positive")
    return options


def main() -> int:
    options=parse_args()
    return rebuild(options) if options.rebuild else run_logger(options)


if __name__ == "__main__":
    raise SystemExit(main())
